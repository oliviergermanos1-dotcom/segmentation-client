#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AGL BUDGET — Intègre les matchs rattrapés dans referentiel_vue + recolore
==========================================================================
Fusionne dans le référentiel les matchs SÛRS retrouvés :
  - orphelins_rattrapes.xlsx : NIVEAU == SUR
  - matchs_valides_claude.xlsx : JUGEMENT ∈ {IDENTIQUE, VALIDE}
(A_REVOIR / PROBABLE écartés.)

Deux entités mono-base reliées par un match retrouvé → fusionnées en une
seule ligne. Les lignes ainsi fusionnées (liens fuzzy/sémantiques, donc à
surveiller) sont COLORÉES en orange. Recalcule ensuite les non-matchés.

Sorties : referentiel_vue_maj.xlsx (coloré) + non_matches_maj.xlsx
"""
import argparse, os, sys
from collections import defaultdict
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

BASES = ["CRM", "STATCOM", "IRIS", "RUBRIKS"]


class UF:
    def __init__(s): s.p = {}
    def find(s, x):
        s.p.setdefault(x, x)
        while s.p[x] != x: s.p[x] = s.p[s.p[x]]; x = s.p[x]
        return x
    def union(s, a, b): s.p[s.find(a)] = s.find(b)


def collect_links(orph_path, valid_path):
    """Retourne liste de (base_a, id_a, base_b, id_b) des matchs sûrs."""
    links = []
    xo = pd.ExcelFile(orph_path)
    for sh in xo.sheet_names:
        if "_x_" not in sh: continue
        ba, bb = sh.split("_x_")
        d = pd.read_excel(xo, sh)
        if "NIVEAU" not in d.columns: continue
        d = d[d["NIVEAU"] == "SUR"]
        for _, r in d.iterrows():
            links.append((ba, str(r[f"ID_{ba}"]), bb, str(r[f"ID_{bb}"])))
    xv = pd.ExcelFile(valid_path)
    for sh in xv.sheet_names:
        if "_x_" not in sh: continue
        ba, bb = sh.split("_x_")
        d = pd.read_excel(xv, sh)
        if "JUGEMENT" not in d.columns: continue
        d = d[d["JUGEMENT"].isin(["IDENTIQUE", "VALIDE"])]
        for _, r in d.iterrows():
            links.append((ba, str(r[f"ID_{ba}"]), bb, str(r[f"ID_{bb}"])))
    return links


def main(argv=None):
    ap = argparse.ArgumentParser(__doc__)
    ap.add_argument("--vue", required=True)
    ap.add_argument("--orphelins", required=True)
    ap.add_argument("--valides", required=True)
    ap.add_argument("--out-vue", default="referentiel_vue_maj.xlsx")
    ap.add_argument("--out-nonmatch", default="non_matches_maj.xlsx")
    args = ap.parse_args(argv)

    v = pd.read_excel(args.vue).fillna("")
    v = v.reset_index(drop=True)
    # index id par base → ligne
    idx = {b: {} for b in BASES}
    for i, row in v.iterrows():
        for b in BASES:
            val = str(row.get(f"ID_{b}", "")).strip()
            if val and val.lower() != "nan":
                idx[b].setdefault(val, i)

    links = collect_links(args.orphelins, args.valides)
    print(f"Liens sûrs à intégrer : {len(links):,}")

    uf = UF(); fused_rows = set(); n_ok = 0
    for ba, ida, bb, idb in links:
        ra, rb = idx[ba].get(ida), idx[bb].get(idb)
        if ra is None or rb is None or ra == rb:
            continue
        uf.union(ra, rb); fused_rows.add(ra); fused_rows.add(rb); n_ok += 1
    print(f"Liens appliqués : {n_ok:,}")

    # regrouper et fusionner
    groups = defaultdict(list)
    for i in range(len(v)):
        groups[uf.find(i)].append(i)

    def pick(members, col):
        for i in members:
            val = str(v.at[i, col]).strip()
            if val and val.lower() != "nan":
                return v.at[i, col]
        return ""

    out_rows = []
    for root, members in groups.items():
        fused = any(m in fused_rows for m in members) and len(members) > 1
        rec = {}
        for b in BASES:
            rec[f"NOM_{b}"] = pick(members, f"NOM_{b}")
            rec[f"ID_{b}"] = pick(members, f"ID_{b}")
        noms = [rec[f"NOM_{b}"] for b in BASES if rec[f"NOM_{b}"]]
        rec["NOM_CONSOLIDE"] = max(noms, key=len) if noms else pick(members, "NOM_CONSOLIDE")
        rec["ID_RETENU"] = (rec["ID_CRM"] or rec["ID_RUBRIKS"] or rec["ID_IRIS"]
                            or rec["ID_STATCOM"] or pick(members, "ID_RETENU"))
        rec["SECTEUR"] = pick(members, "SECTEUR")
        rec["CONFIANCE_LIEN"] = "RATTRAPE" if fused else pick(members, "CONFIANCE_LIEN")
        rec["_fused"] = fused
        out_rows.append(rec)

    cols = []
    for b in BASES: cols += [f"NOM_{b}", f"ID_{b}"]
    cols += ["NOM_CONSOLIDE", "ID_RETENU", "SECTEUR", "CONFIANCE_LIEN"]
    out = pd.DataFrame(out_rows)
    out["_nb"] = (out[[f"ID_{b}" for b in BASES]] != "").sum(axis=1)
    out = out.sort_values(["_fused", "_nb"], ascending=False)
    fused_flag = out["_fused"].tolist()
    out = out[cols]

    print(f"Entités : {len(v):,} → {len(out):,} (−{len(v)-len(out):,} fusions)")
    out.to_excel(args.out_vue, index=False)

    # colorer les lignes rattrapées (orange clair)
    wb = load_workbook(args.out_vue); ws = wb.active
    fill = PatternFill("solid", fgColor="FFE0B2")
    for r, isf in enumerate(fused_flag, start=2):
        if isf:
            for c in range(1, len(cols) + 1):
                ws.cell(row=r, column=c).fill = fill
    wb.save(args.out_vue)
    print(f"[ok] {sum(fused_flag):,} lignes colorées (rattrapées) → {args.out_vue}")

    # non-matchés mis à jour
    pres = {b: out[f"NOM_{b}"].astype(str).str.strip() != "" for b in BASES}
    nb = sum(pres[b].astype(int) for b in BASES)
    recap = []
    with pd.ExcelWriter(args.out_nonmatch, engine="openpyxl") as w:
        for b in BASES:
            sub = out[pres[b] & (nb == 1)][[f"NOM_{b}", f"ID_{b}", "SECTEUR"]].copy()
            sub.columns = ["NOM", "ID", "SECTEUR"]
            sub.sort_values("NOM").to_excel(w, sheet_name=f"non_match_{b}", index=False)
            recap.append({"BASE": b, "NON_MATCHEES": len(sub)})
            print(f"  {b}: {len(sub):,} non matchées (après intégration)")
        pd.DataFrame(recap).to_excel(w, sheet_name="synthese", index=False)
    print(f"[ok] → {args.out_nonmatch}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
